import sys
try:
    from pptx import Presentation
except ImportError:
    print("python-pptx not installed yet")
    sys.exit(1)

prs = Presentation("Scaling_Test-Time_Compute_via_Kolmogorov-Arnold_Energy_Models.pptx")
with open("pptx_dump.txt", "w", encoding="utf-8") as f:
    for i, slide in enumerate(prs.slides):
        f.write(f"\\n--- SLIDE {i+1} ---\\n")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                f.write(shape.text + "\\n")
print("Dumped to pptx_dump.txt")
